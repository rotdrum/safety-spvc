import csv
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


INPUT = Path("/Users/kkiattikun/Downloads/89121_26052213133250.xlsx")
OUTPUT_DIR = Path("/Users/kkiattikun/Documents/Patthama-Fanz/student-late-system/outputs/import_students")

STUDENT_HEADERS = [
    "student_id",
    "student_no",
    "prefix",
    "first_name",
    "last_name",
    "full_name",
    "grade",
    "room",
    "homeroom_teacher",
    "status",
    "card_code",
    "notes",
]

PREFIXES = ["นางสาว", "เด็กหญิง", "เด็กชาย", "น.ส.", "นาย", "นาง"]


def clean(value):
    if value is None:
        return ""
    text = str(value).strip()
    return re.sub(r"\s+", " ", text)


def as_int_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return clean(value)


def split_name(full_name):
    name = clean(full_name)
    prefix = ""
    rest = name
    for candidate in PREFIXES:
        if name.startswith(candidate):
            prefix = candidate
            rest = name[len(candidate):].strip()
            break

    parts = rest.split()
    first_name = parts[0] if parts else ""
    last_name = " ".join(parts[1:]) if len(parts) > 1 else ""
    return prefix, first_name, last_name, " ".join([prefix + first_name if prefix else first_name, last_name]).strip()


def level_from_class(class_year):
    text = clean(class_year)
    if "/" in text:
        return text.split("/", 1)[0]
    return text


def find_header_row(ws):
    for row_idx in range(1, min(ws.max_row, 30) + 1):
        values = [clean(ws.cell(row_idx, col).value) for col in range(1, min(ws.max_column, 10) + 1)]
        if "รหัสประจำตัว" in values and "ชื่อ - สกุล" in values:
            return row_idx, values
    return None, []


def col_index(headers, name):
    try:
        return headers.index(name) + 1
    except ValueError:
        return 0


def main():
    wb = load_workbook(INPUT, read_only=False, data_only=True)
    rows = []
    rooms = []
    room_seen = set()
    sheets = []
    status_counts = {}
    duplicate_ids = {}

    for ws in wb.worksheets:
        header_row, headers = find_header_row(ws)
        if not header_row:
            sheets.append({"sheet": ws.title, "students": 0, "skipped": "no student header"})
            continue

        group_code = clean(ws["C7"].value)
        group_name = clean(ws["E7"].value)
        class_year = clean(ws["C8"].value)
        advisor = clean(ws["E8"].value)
        grade = level_from_class(class_year)
        room = group_name or class_year or ws.title
        if room and room not in room_seen:
            room_seen.add(room)
            rooms.append(room)

        no_col = col_index(headers, "ลำดับ")
        student_id_col = col_index(headers, "รหัสประจำตัว")
        full_name_col = col_index(headers, "ชื่อ - สกุล")
        learner_type_col = col_index(headers, "ประเภทผู้เรียน")
        status_col = col_index(headers, "สถานะนักเรียนนักศึกษา")

        count = 0
        for row_idx in range(header_row + 1, ws.max_row + 1):
            student_id = as_int_text(ws.cell(row_idx, student_id_col).value)
            full_name_raw = clean(ws.cell(row_idx, full_name_col).value)
            if not student_id and not full_name_raw:
                continue
            if not student_id or not full_name_raw:
                continue

            student_no = as_int_text(ws.cell(row_idx, no_col).value)
            learner_type = clean(ws.cell(row_idx, learner_type_col).value)
            source_status = clean(ws.cell(row_idx, status_col).value)
            status = "Active" if source_status == "กำลังศึกษา" else "Inactive"
            prefix, first_name, last_name, full_name = split_name(full_name_raw)
            notes = " | ".join(
                part
                for part in [
                    f"sheet={ws.title}",
                    f"group_code={group_code}" if group_code else "",
                    f"class={class_year}" if class_year else "",
                    f"learner_type={learner_type}" if learner_type else "",
                    f"source_status={source_status}" if source_status else "",
                ]
                if part
            )

            rows.append(
                [
                    student_id,
                    student_no,
                    prefix,
                    first_name,
                    last_name,
                    full_name,
                    grade,
                    room,
                    advisor,
                    status,
                    "",
                    notes,
                ]
            )
            duplicate_ids.setdefault(student_id, 0)
            duplicate_ids[student_id] += 1
            status_counts[source_status or "(blank)"] = status_counts.get(source_status or "(blank)", 0) + 1
            count += 1

        sheets.append({"sheet": ws.title, "students": count, "room": room, "advisor": advisor})

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    students_tsv = OUTPUT_DIR / "students.tsv"
    with students_tsv.open("w", encoding="utf-8", newline="") as file:
        writer = csv.writer(file, delimiter="\t", lineterminator="\n")
        writer.writerow(STUDENT_HEADERS)
        writer.writerows(rows)

    rooms_tsv = OUTPUT_DIR / "rooms.tsv"
    with rooms_tsv.open("w", encoding="utf-8", newline="") as file:
        writer = csv.writer(file, delimiter="\t", lineterminator="\n")
        writer.writerow(["Rooms"])
        for room in rooms:
            writer.writerow([room])

    summary = {
        "workbook": str(INPUT),
        "sheet_count": len(wb.worksheets),
        "student_count": len(rows),
        "room_count": len(rooms),
        "status_counts": status_counts,
        "duplicate_student_ids": {key: value for key, value in duplicate_ids.items() if value > 1},
        "sheets": sheets,
        "outputs": {
            "students_tsv": str(students_tsv),
            "rooms_tsv": str(rooms_tsv),
        },
    }
    (OUTPUT_DIR / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    sys.exit(main())
